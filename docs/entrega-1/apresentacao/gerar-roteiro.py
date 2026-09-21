"""Gera roteiro-apresentacao.pdf a partir das notas do apresentador do PPTX.
Uso: pip install python-pptx reportlab && python gerar-roteiro.py
"""
import os
from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PASTA = os.path.dirname(os.path.abspath(__file__))
MARROM, TOMATE, CINZA = colors.HexColor("#2B1B17"), colors.HexColor("#C8352B"), colors.HexColor("#6B5B55")

# (slides, quem fala, área, tempo sugerido, o que mostrar na demonstração)
PARTES = [
    ([1], "Yuri Mascarenhas", "Abertura", "1 min", None),
    ([2, 3], "Yuri Mascarenhas", "Autenticação, sessões e controle de acesso", "3 min",
     "Abrir ValidadorCookieLimitesTest e rodar: mvn test -Dtest=ValidadorCookie*"),
    ([4, 5], "Vinicius Rocha Pinheiro", "Carrinho, pedidos e checkout", "3 min",
     "Rodar: mvn test -Pdefeitos -Dtest=ComprarRegrasTest e mostrar 'expected: <35.0> but was: <20.0>'"),
    ([6, 7], "Felipe Paixão", "Gestão de clientes e endereços", "3 min",
     "Abrir DaoClienteTest (deveRecusarClienteInativo) e mostrar em cadastro.js a condição !field.name == 'complemento'"),
    ([8, 9], "Vinicius Fonseca de Freitas", "Cardápio e montagem de lanches", "3 min",
     "Mostrar os dois laços while(keys.hasNext()) em salvarLancheCliente.java e rodar: mvn test -Pdefeitos -Dtest=SalvarLancheClienteTest"),
    ([10, 11], "Thiago Ferreira", "Gestão de estoque e insumos", "3 min",
     "Abrir IngredienteControllersTest (deveBloquearQuemNaoEhFuncionario e caracterizacaoValoresNegativosSaoAceitos)"),
    ([12], "Thiago Ferreira", "Encerramento", "1 min", "Rodar mvn clean test e mostrar 'Tests run: 115, Failures: 0'"),
]

PERGUNTAS = [
    ("Por que alguns testes ficam vermelhos de propósito?",
     "Eles descrevem o comportamento correto. Enquanto o defeito existir, falham. Ficam no perfil -Pdefeitos para não quebrar o build normal, e não usamos @Disabled para não esconder o problema."),
    ("Vocês corrigiram os defeitos?", "Não. O objetivo da entrega era testar e reportar; cada defeito tem um teste que o reproduz."),
    ("Por que mexeram no código de produção?",
     "Só adicionamos construtores que recebem a conexão ou os DAOs, para poder usar mocks. O construtor padrão e o comportamento no Tomcat não mudaram."),
    ("A classe da área de estoque não é CRUD?",
     "DaoIngrediente é CRUD, e assumimos isso. O sistema não tem regra de estoque; usamos os servlets de insumo como melhor alternativa e registramos a lacuna."),
    ("O que a cobertura garante?", "Só que aquelas linhas e ramos foram executados pelos testes, não que não existam defeitos."),
]


def estilos():
    base = dict(fontName="Helvetica", textColor=MARROM, alignment=TA_LEFT)
    return {
        "titulo": ParagraphStyle("titulo", fontSize=22, leading=26, fontName="Helvetica-Bold", textColor=MARROM),
        "sub": ParagraphStyle("sub", fontSize=10.5, leading=14, textColor=CINZA, fontName="Helvetica"),
        "quem": ParagraphStyle("quem", fontSize=15, leading=19, fontName="Helvetica-Bold", textColor=MARROM, spaceBefore=4),
        "rotulo": ParagraphStyle("rotulo", fontSize=8.5, leading=11, fontName="Helvetica-Bold", textColor=TOMATE, spaceBefore=8),
        "fala": ParagraphStyle("fala", fontSize=11.5, leading=16.5, **base),
        "nota": ParagraphStyle("nota", fontSize=10, leading=14, **base),
    }


def main():
    deck = Presentation(os.path.join(PASTA, "apresentacao-entrega-1.pptx"))
    notas = {i: s.notes_slide.notes_text_frame.text.strip() for i, s in enumerate(deck.slides, 1)}
    e = estilos()
    doc = SimpleDocTemplate(os.path.join(PASTA, "roteiro-apresentacao.pdf"), pagesize=A4, title="Roteiro - Entrega 1",
                            leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    h = [Paragraph("Roteiro da apresentação", e["titulo"]),
         Paragraph("Code Burguer's · Qualidade e Teste de Software · Entrega 1 · 12 slides, cerca de 17 minutos", e["sub"]),
         Spacer(1, 10)]

    linhas = [["Slides", "Quem fala", "Parte", "Tempo"]]
    for slides, quem, area, tempo, _ in PARTES:
        linhas.append(["-".join(str(n) for n in slides) if len(slides) == 1 else f"{slides[0]} e {slides[1]}", quem, area, tempo])
    tabela = Table(linhas, colWidths=[1.8 * cm, 5.2 * cm, 8.0 * cm, 2.0 * cm])
    tabela.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), MARROM), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F1EE")]),
        ("TEXTCOLOR", (0, 1), (-1, -1), MARROM), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    h += [tabela, Spacer(1, 8),
          Paragraph("As falas abaixo são as mesmas das notas do apresentador no PowerPoint. Não precisa decorar: use como guia "
                    "e fale com suas palavras. Quem apresenta deve conseguir abrir o teste citado e explicar o que ele verifica.", e["nota"]),
          PageBreak()]

    for slides, quem, area, tempo, demonstracao in PARTES:
        bloco = [Paragraph(quem, e["quem"]), Paragraph(f"{area} · {tempo}", e["sub"])]
        for n in slides:
            bloco += [Paragraph(f"SLIDE {n}", e["rotulo"]), Paragraph(notas[n], e["fala"])]
        if demonstracao:
            bloco += [Paragraph("PARA MOSTRAR", e["rotulo"]), Paragraph(demonstracao, e["nota"])]
        bloco.append(Spacer(1, 16))
        h.append(KeepTogether(bloco))

    perguntas = [Paragraph("Perguntas prováveis", e["quem"]), Paragraph("Qualquer integrante pode responder", e["sub"])]
    for pergunta, resposta in PERGUNTAS:
        perguntas += [Paragraph(pergunta.upper(), e["rotulo"]), Paragraph(resposta, e["nota"])]
    h.append(KeepTogether(perguntas))

    doc.build(h)
    print("gerado: roteiro-apresentacao.pdf")


if __name__ == "__main__":
    main()
